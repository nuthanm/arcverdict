"""Export the workbook outputs the equity desk renders."""

import json
from pathlib import Path

import openpyxl

DOWNLOADS = Path(r"c:\Users\NuthanMurarysetty\Downloads")
OUT = Path(__file__).resolve().parents[1] / "data"
DAILY_TAIL = 20


def num(value, places=2):
    if value is None or isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return round(float(value), places)
    return value


def day(value):
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return None


def header_index(header):
    return {name: i for i, name in enumerate(header) if name}


def main():
    OUT.mkdir(exist_ok=True)

    metrics = {}
    wb = openpyxl.load_workbook(DOWNLOADS / "History_Data_Metric.xlsx", read_only=True, data_only=True)
    ws = wb["Summary"]
    rows = ws.iter_rows(values_only=True)
    header = header_index(next(rows))
    for row in rows:
        company = row[header["Company"]]
        if not company:
            continue
        metrics[company] = {
            "trend": row[header["Trend"]],
            "trendScore": num(row[header["Trend Score"]], 0),
            "totalScore": num(row[header["Total Score"]]),
            "runningSum": num(row[header["Running Sum"]]),
            "runningSumPercentile": num(row[header["Running Sum Percentile"]]),
            "slopeStrength": num(row[header["Slope Strength"]]),
            "consistency": num(row[header["Consistency"]]),
            "momentum": num(row[header["Momentum"]]),
            "acceleration": num(row[header["Acceleration"]]),
            "ema20": num(row[header["EMA20"]]),
            "ema50": num(row[header["EMA50"]]),
            "ema200": num(row[header["EMA200"]]),
            "rsi": num(row[header["Wilder_RSI14"]]),
            "adx": num(row[header["ADX"]]),
            "atrPct": num(row[header["ATR14_Pct"]]),
            "bbPosition": num(row[header["BB_Position"]], 3),
        }
    wb.close()

    research = {}
    daily = {}
    current = None
    tail = []
    wb = openpyxl.load_workbook(DOWNLOADS / "Trend_Lifecycle_Data.xlsx", read_only=True, data_only=True)

    ws = wb["Research Summary"]
    rows = ws.iter_rows(values_only=True)
    header = header_index(next(rows))
    for row in rows:
        company = row[header["Company"]]
        if not company:
            continue
        research[company] = {
            "rows": row[header["Rows"]],
            "startDate": day(row[header["Start Date"]]),
            "endDate": day(row[header["End Date"]]),
            "candidateRows": row[header["Candidate Rows"]],
            "confirmedDecayRows": row[header["Confirmed Decay Rows"]],
        }

    ws = wb["Daily Lifecycle"]
    rows = ws.iter_rows(values_only=True)
    header = header_index(next(rows))

    def flush():
        if current is not None:
            daily[current] = tail[-DAILY_TAIL:]

    for row in rows:
        company = row[header["Company"]]
        if not company:
            continue
        if company != current:
            flush()
            current = company
            tail = []
        tail.append(
            {
                "date": day(row[header["Date"]]),
                "close": num(row[header["Close"]]),
                "trend": row[header["Trend"]],
                "trendScore": num(row[header["Trend Score"]], 0),
                "healthScore": num(row[header["Trend Health Score"]]),
                "healthState": row[header["Health State"]],
                "pressure": row[header["Pressure"]],
                "structureState": row[header["Structure State"]],
                "decayStatus": row[header["Decay Status"]],
                "candidateStatus": row[header["Candidate Status"]],
                "suggestion": row[header["Suggestion"]],
            }
        )
    flush()

    latest = []
    ws = wb["Latest Lifecycle"]
    rows = ws.iter_rows(values_only=True)
    header = header_index(next(rows))
    for row in rows:
        company = row[header["Company"]]
        if not company:
            continue
        latest.append(
            {
                "company": company,
                "date": day(row[header["Date"]]),
                "close": num(row[header["Close"]]),
                "healthScore": num(row[header["Trend Health Score"]]),
                "healthState": row[header["Health State"]],
                "pressure": row[header["Pressure"]],
                "healthPressure1d": num(row[header["Health Pressure 1D"]]),
                "healthPressure5d": num(row[header["Health Pressure 5D"]]),
                "slopeStrengthLevel": num(row[header["Slope Strength Level"]]),
                "runningSumLevel": num(row[header["Running Sum Level"]]),
                "rsiLevel": num(row[header["Wilder RSI14 Level"]]),
                "volumeParticipation": num(row[header["Volume Participation"]], 3),
                "volumeParticipationLevel": num(row[header["Volume Participation Level"]]),
                "deterioratingComponents": row[header["Deteriorating Components"]],
                "improvingComponents": row[header["Improving Components"]],
                "deteriorationPersistence": row[header["Deterioration Persistence"]],
                "priceSlope3d": num(row[header["Price Slope 3D"]]),
                "lowerHigh": bool(row[header["Lower High"]]) if row[header["Lower High"]] is not None else None,
                "lowerLow": bool(row[header["Lower Low"]]) if row[header["Lower Low"]] is not None else None,
                "structureState": row[header["Structure State"]],
                "decayStatus": row[header["Decay Status"]],
                "candidateStatus": row[header["Candidate Status"]],
                "suggestion": row[header["Suggestion"]],
                "metrics": metrics.get(company),
                "research": research.get(company),
            }
        )
    wb.close()

    latest.sort(key=lambda row: (row["healthScore"] is None, row["healthScore"] if row["healthScore"] is not None else 0))

    (OUT / "lifecycle-latest.json").write_text(json.dumps(latest), encoding="utf-8")
    (OUT / "lifecycle-daily.json").write_text(json.dumps(daily), encoding="utf-8")
    print("latest", len(latest), "daily", len(daily))
    print("latest bytes", (OUT / "lifecycle-latest.json").stat().st_size)
    print("daily bytes", (OUT / "lifecycle-daily.json").stat().st_size)


if __name__ == "__main__":
    main()
